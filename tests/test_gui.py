"""Headless UI integration, actual Qt widgets, no model and no private documents."""
import json
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from localdoc.qt_compat import QApplication, QFileDialog, QMessageBox, QPushButton
from localdoc.gui import Window, Editor
from localdoc.core import FIELDS, Record, fingerprint
from localdoc.qt_compat import QT_BINDING
if QT_BINDING == 'PySide2':
    from shiboken2 import delete as delete_qt
else:
    from shiboken6 import delete as delete_qt


class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.app = QApplication.instance() or QApplication([])
    @classmethod
    def tearDownClass(cls):
        # Destroy native widgets while Python and the binding are still alive.
        # Letting Qt5 tear them down during interpreter finalization can crash.
        cls.app.processEvents()
        delete_qt(cls.app)
        cls.app = None
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.state_patch = patch('localdoc.gui.state_root', return_value=self.root)
        self.state_patch.start()
        self.window = Window()
        source = self.root/'source.txt'; source.write_text('Договор 125 ООО Альфа 15.10.2026', encoding='utf-8')
        self.data = {k: None for k in FIELDS}
        self.data.update(document_type='Договор', organization='ООО Альфа', document_number='125', document_date='2026-10-15')
        self.data['evidence'] = {k: None for k in FIELDS}
        self.record = Record(str(source), fingerprint(source), self.data, text=source.read_text('utf-8'))
        self.window.records = [self.record]
    def tearDown(self):
        self.window.close()
        delete_qt(self.window)
        self.app.processEvents()
        self.state_patch.stop(); self.tmp.cleanup()
    def test_manual_review_and_preview(self):
        editor = Editor(self.record, self.window)
        editor.make_name()
        self.assertIn('2026-10-15_Договор_125_ООО_Альфа', editor.name.text())
        editor.confirm.setChecked(True)
        editor.save()
        self.assertTrue(self.record.approved)
        self.window.refresh()
        self.assertEqual(self.window.table.item(0,2).text(), 'Подтвержден')
    def test_export_json(self):
        target = self.root/'export.json'
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(target),'')), patch.object(QMessageBox, 'information'):
            self.window.export('json')
        result = json.loads(target.read_text('utf-8'))
        self.assertIsNone(result[0]['inn'])
        self.assertIn('inn', result[0]['review_required'])
    def test_xlsx_formula_injection_prevented(self):
        from openpyxl import load_workbook
        target = self.root/'export.xlsx'
        self.record.data['subject'] = '=HYPERLINK("https://example.invalid","x")'
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(target),'')), patch.object(QMessageBox, 'information'):
            self.window.export('xlsx')
        book = load_workbook(target, data_only=False)
        try:
            cell = book.active.cell(2, 6) # source + type/organization/number/date/subject
            self.assertEqual(cell.data_type, 's')
            self.assertTrue(cell.value.startswith('='))
        finally: book.close()
    def test_block_export_over_original_json_input(self):
        # Protection is also applied to paths retained by the rename journal.
        source = self.root/'original.json'; source.write_text('original')
        self.record.source = str(source)
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(source),'')), patch.object(QMessageBox, 'warning') as warning:
            self.window.export('json')
        self.assertTrue(warning.called); self.assertEqual(source.read_text(), 'original')
    def test_error_cannot_be_manually_approved(self):
        self.record.error = 'OCR error'
        editor = Editor(self.record, self.window)
        save = [b for b in editor.findChildren(QPushButton) if b.text() == 'Сохранить проверку'][0]
        self.assertFalse(save.isEnabled())
